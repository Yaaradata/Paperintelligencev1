#!/usr/bin/env python3
"""Export blind double-label workbook for labeller=ranjith, round=v1.

30 rows sampled across strata and across Subha's verdict values.
Same blind format as golden_labelling_200.xlsx (no model scores).

  PYTHONPATH=src python3 scripts/export_double_label_30.py
"""

from __future__ import annotations

import json
import random
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

import yaml  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from paper_intelligence.db import connect  # noqa: E402

GOLDEN_DIR = ROOT / "golden"
XLSX_OUT = GOLDEN_DIR / "double_label_30.xlsx"
META_OUT = GOLDEN_DIR / "double_label_30_sample.json"
CLASS_POL = ROOT / "policies" / "classification" / "v002.yaml"
SEATS_POL = ROOT / "policies" / "editorial_seats" / "v001.yaml"
QUALITY_PROMPT = ROOT / "prompts" / "quality" / "v002.md"

TARGET_N = 30
SEED = 42
VERDICTS = ("winner_material", "shortlist", "maybe", "reject")

SCORE_COLS = (
    "h_ai_relevance",
    "h_technical_significance",
    "h_apparent_novelty",
    "h_practical_applicability",
    "h_professional_value",
    "h_learning_value",
    "h_evidence_strength",
    "h_tech_relevance",
    "h_product_relevance",
    "h_final_score",
)

LABEL_HEADERS = (
    "row_no",
    "paper_id",
    "arxiv_id",
    "arxiv_url",
    "title",
    "abstract",
    "published_at",
    "organisations",
    "sample_stratum",
    *SCORE_COLS,
    "h_domain",
    "h_application_domain",
    "h_newsletter_verdict",
    "h_reject_reason",
    "h_notes",
)


def _half_steps() -> list[str]:
    return [f"{i / 2:.1f}" for i in range(0, 21)]


def _load_vocab() -> tuple[list[str], list[str]]:
    data = yaml.safe_load(CLASS_POL.read_text(encoding="utf-8")) or {}
    domains = [str(x) for x in (data.get("domain") or [])]
    apps = [str(x) for x in (data.get("application_domain") or [])]
    if not domains or not apps:
        raise RuntimeError(f"empty domain vocab in {CLASS_POL}")
    return domains, apps


def _sample_30(rng: random.Random) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Stratify by (sample_stratum × verdict), then fill to 30."""
    with connect() as conn:
        humans = conn.execute(
            """
            SELECT g.paper_id, g.arxiv_id, g.sample_stratum, g.h_newsletter_verdict,
                   g.h_final_score,
                   p.title, p.abstract, p.published_at
            FROM paper_intelligence.golden_human_scores g
            JOIN paper_intelligence.papers p ON p.paper_id = g.paper_id
            WHERE g.labeller = 'subha' AND g.label_round = 'v1'
            ORDER BY g.paper_id
            """
        ).fetchall()
        ids = [int(r["paper_id"]) for r in humans]
        orgs = conn.execute(
            """
            SELECT content_item_id, array_agg(DISTINCT organisation_name ORDER BY organisation_name) AS orgs
            FROM paper_intelligence.paper_organisations
            WHERE content_item_id = ANY(%s)
            GROUP BY content_item_id
            """,
            (ids,),
        ).fetchall() if False else []
        # organisations table name may vary — try soft lookup
        org_map: dict[int, str] = {}
        try:
            orows = conn.execute(
                """
                SELECT paper_id, string_agg(DISTINCT org_name, '; ' ORDER BY org_name) AS orgs
                FROM paper_intelligence.paper_org_mentions
                WHERE paper_id = ANY(%s)
                GROUP BY paper_id
                """,
                (ids,),
            ).fetchall()
            org_map = {int(r["paper_id"]): r["orgs"] or "" for r in orows}
        except Exception:
            try:
                conn.rollback()
            except Exception:
                pass
            try:
                orows = conn.execute(
                    """
                    SELECT content_item_id AS paper_id,
                           string_agg(DISTINCT organisation, '; ') AS orgs
                    FROM paper_intelligence.author_affiliations
                    WHERE content_item_id = ANY(%s)
                    GROUP BY content_item_id
                    """,
                    (ids,),
                ).fetchall()
                org_map = {int(r["paper_id"]): r["orgs"] or "" for r in orows}
            except Exception:
                try:
                    conn.rollback()
                except Exception:
                    pass

    pool = []
    for r in humans:
        aid = r.get("arxiv_id") or ""
        pool.append(
            {
                "paper_id": int(r["paper_id"]),
                "arxiv_id": aid,
                "arxiv_url": f"https://arxiv.org/abs/{aid}" if aid else "",
                "title": r.get("title") or "",
                "abstract": r.get("abstract") or "",
                "published_at": str(r["published_at"]) if r.get("published_at") else "",
                "organisations": org_map.get(int(r["paper_id"]), ""),
                "sample_stratum": r.get("sample_stratum") or "unknown",
                "subha_verdict": r.get("h_newsletter_verdict") or "unknown",
                "subha_final": float(r["h_final_score"])
                if r.get("h_final_score") is not None
                else None,
            }
        )

    by_cell: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in pool:
        by_cell[(row["sample_stratum"], row["subha_verdict"])].append(row)

    # Target: cover every non-empty cell at least once, then pad by verdict balance
    selected: list[dict[str, Any]] = []
    seen: set[int] = set()

    cells = sorted(by_cell.keys())
    rng.shuffle(cells)
    for cell in cells:
        candidates = list(by_cell[cell])
        rng.shuffle(candidates)
        pick = candidates[0]
        if pick["paper_id"] not in seen:
            selected.append(pick)
            seen.add(pick["paper_id"])

    # Pad to TARGET_N with remaining rows, balancing verdicts
    remaining = [r for r in pool if r["paper_id"] not in seen]
    rng.shuffle(remaining)

    def verdict_counts(rows: list[dict[str, Any]]) -> dict[str, int]:
        c: dict[str, int] = {v: 0 for v in VERDICTS}
        for r in rows:
            v = r["subha_verdict"]
            if v in c:
                c[v] += 1
        return c

    while len(selected) < TARGET_N and remaining:
        counts = verdict_counts(selected)
        # Prefer under-represented verdicts
        remaining.sort(
            key=lambda r: (counts.get(r["subha_verdict"], 0), rng.random())
        )
        pick = remaining.pop(0)
        selected.append(pick)
        seen.add(pick["paper_id"])

    # Stable display order: stratum then paper_id (blind — no verdict column shown)
    selected.sort(key=lambda r: (r["sample_stratum"], r["paper_id"]))

    meta = {
        "seed": SEED,
        "n": len(selected),
        "labeller_target": "ranjith",
        "label_round": "v1",
        "stratum_counts": dict(
            __import__("collections").Counter(r["sample_stratum"] for r in selected)
        ),
        "subha_verdict_counts": dict(
            __import__("collections").Counter(r["subha_verdict"] for r in selected)
        ),
        # Hidden from workbook — for ceiling analysis after import only
        "sample_paper_ids": [r["paper_id"] for r in selected],
        "subha_labels_hidden": [
            {
                "paper_id": r["paper_id"],
                "stratum": r["sample_stratum"],
                "verdict": r["subha_verdict"],
                "h_final": r["subha_final"],
            }
            for r in selected
        ],
    }
    return selected, meta


def _write_xlsx(rows: list[dict[str, Any]], domains: list[str], apps: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "labelling"
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    for col, name in enumerate(LABEL_HEADERS, start=1):
        cell = ws.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, row in enumerate(rows, start=1):
        values = {
            "row_no": i,
            "paper_id": row["paper_id"],
            "arxiv_id": row["arxiv_id"],
            "arxiv_url": row["arxiv_url"],
            "title": row["title"],
            "abstract": row["abstract"],
            "published_at": row["published_at"],
            "organisations": row["organisations"],
            "sample_stratum": row["sample_stratum"],
        }
        for col, name in enumerate(LABEL_HEADERS, start=1):
            val = values.get(name, "")
            cell = ws.cell(i + 1, col, val if val is not None else "")
            if name in {"title", "abstract", "organisations"}:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "D2"
    widths = {
        "A": 8,
        "B": 12,
        "C": 14,
        "D": 28,
        "E": 40,
        "F": 60,
        "G": 20,
        "H": 28,
        "I": 16,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for col in range(10, len(LABEL_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.row_dimensions[1].height = 30
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 60

    ws_lists = wb.create_sheet("_lists")
    for i, d in enumerate(domains, start=1):
        ws_lists.cell(i, 1, d)
    for i, a in enumerate(apps, start=1):
        ws_lists.cell(i, 2, a)
    for i, v in enumerate(VERDICTS, start=1):
        ws_lists.cell(i, 3, v)
    for i, s in enumerate(_half_steps(), start=1):
        ws_lists.cell(i, 4, float(s))
    ws_lists.sheet_state = "hidden"

    score_start = LABEL_HEADERS.index("h_ai_relevance") + 1
    score_end = LABEL_HEADERS.index("h_final_score") + 1
    dv_score = DataValidation(
        type="list",
        formula1=f"=_lists!$D$1:$D${len(_half_steps())}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid score",
        error="Enter 0.0–10.0 in 0.5 steps (use the dropdown).",
        showDropDown=False,
    )
    dv_score.add(
        f"{get_column_letter(score_start)}2:{get_column_letter(score_end)}{len(rows)+1}"
    )
    ws.add_data_validation(dv_score)

    dom_col = LABEL_HEADERS.index("h_domain") + 1
    app_col = LABEL_HEADERS.index("h_application_domain") + 1
    verd_col = LABEL_HEADERS.index("h_newsletter_verdict") + 1
    for col, formula, title, err in (
        (dom_col, f"=_lists!$A$1:$A${len(domains)}", "Invalid domain", "Pick closed vocab."),
        (app_col, f"=_lists!$B$1:$B${len(apps)}", "Invalid application_domain", "Pick closed vocab."),
        (verd_col, f"=_lists!$C$1:$C${len(VERDICTS)}", "Invalid verdict", "winner_material/shortlist/maybe/reject"),
    ):
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle=title,
            error=err,
        )
        dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{len(rows)+1}")
        ws.add_data_validation(dv)

    wr = wb.create_sheet("rubric", 1)
    seats = yaml.safe_load(SEATS_POL.read_text(encoding="utf-8")) or {}
    quality_text = QUALITY_PROMPT.read_text(encoding="utf-8")
    wr["A1"] = "Quality rubric (six dimensions + composite)"
    wr["A1"].font = Font(bold=True, size=14)
    wr["A3"] = quality_text
    wr["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    wr.merge_cells("A3:D40")
    wr["A42"] = "TECH and PRODUCT seat definitions"
    wr["A42"].font = Font(bold=True, size=14)
    wr["A44"] = yaml.dump(seats.get("seats") or seats, sort_keys=False, allow_unicode=True)
    wr["A44"].alignment = Alignment(wrap_text=True, vertical="top")
    wr.merge_cells("A44:D80")
    wr.column_dimensions["A"].width = 100

    wi = wb.create_sheet("instructions", 2)
    instructions = [
        "Double-label instructions (labeller=ranjith, round=v1)",
        "",
        "1. Label BLIND. Do not look at Subha's labels or any Terra/Jev scores.",
        "2. This is a 30-paper reliability set sampled across strata and verdict mix.",
        "3. Score numeric columns 0–10 in 0.5 steps via dropdowns.",
        "4. h_final_score is YOUR overall newsletter judgement — not a formula.",
        "5. h_newsletter_verdict: winner_material | shortlist | maybe | reject.",
        "6. Leave a row blank rather than guessing.",
        "7. Save and return for import:",
        "     PYTHONPATH=src python3 scripts/import_golden_labels.py \\",
        "       --xlsx golden/double_label_30.xlsx --labeller ranjith --label-round v1",
        "",
        "Purpose: human–human ceiling before judging whether Spearman ~0.6 is good.",
    ]
    for i, line in enumerate(instructions, start=1):
        wi.cell(i, 1, line)
        if i == 1:
            wi.cell(i, 1).font = Font(bold=True, size=14)
    wi.column_dimensions["A"].width = 110

    forbidden = ("terra", "jev", "subha", "rank_gap", "composite")
    for cell in ws[1]:
        name = str(cell.value or "").lower()
        if any(tok in name for tok in forbidden):
            raise RuntimeError(f"labelling sheet leaked column: {cell.value}")

    GOLDEN_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_OUT)


def main() -> int:
    domains, apps = _load_vocab()
    rng = random.Random(SEED)
    rows, meta = _sample_30(rng)
    if len(rows) != TARGET_N:
        print(f"WARNING: got {len(rows)} rows (target {TARGET_N})", flush=True)
    _write_xlsx(rows, domains, apps)
    META_OUT.write_text(json.dumps(meta, indent=2), encoding="utf-8")
    print(f"wrote {XLSX_OUT} n={len(rows)}", flush=True)
    print(f"wrote {META_OUT}", flush=True)
    print(json.dumps({k: meta[k] for k in ("stratum_counts", "subha_verdict_counts")}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
