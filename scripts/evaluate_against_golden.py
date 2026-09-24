#!/usr/bin/env python3
"""Phase G3a — evaluate Terra / Jev engines against human golden labels.

Read-only. No paid calls. Does not change models, prompts, policies, or thresholds.

  PYTHONPATH=src python3 scripts/evaluate_against_golden.py \\
    --labeller subha --label-round v1

Writes:
  reports/golden/engine_comparison.md
  reports/golden/engine_comparison.json
  reports/golden/baseline_v1.md
  golden/golden_scores.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import statistics
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from paper_intelligence.db import connect  # noqa: E402

REPORT_DIR = ROOT / "reports" / "golden"
MD_OUT = REPORT_DIR / "engine_comparison.md"
JSON_OUT = REPORT_DIR / "engine_comparison.json"
BASELINE_OUT = REPORT_DIR / "baseline_v1.md"
XLSX_OUT = ROOT / "golden" / "golden_scores.xlsx"
JEV_V001 = ROOT / "reports" / "review_fixes" / "jev_j1b_quality_scores.json"
JEV_V002 = ROOT / "reports" / "review_fixes" / "jev_j1b_quality_scores_v002.json"

QUALITY_DIMS = (
    "technical_significance",
    "apparent_novelty",
    "practical_applicability",
    "professional_value",
    "learning_value",
    "evidence_strength",
)
Q_WEIGHTS = {
    "technical_significance": 0.28,
    "apparent_novelty": 0.24,
    "practical_applicability": 0.20,
    "professional_value": 0.16,
    "learning_value": 0.12,
}
ENGINES = ("terra", "jev_v001", "jev_v002")
MIN_CELL_N = 25
G3B_RUN_META = REPORT_DIR / "g3b_run.json"


def _excluded_terra_run_ids() -> list[str]:
    """G3b eval run_ids must not pollute the main comparison."""
    if not G3B_RUN_META.exists():
        return []
    try:
        meta = json.loads(G3B_RUN_META.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    rid = meta.get("run_id")
    return [str(rid)] if rid else []


def _args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--labeller", default="subha")
    p.add_argument("--label-round", default="v1")
    p.add_argument("--min-cell-n", type=int, default=MIN_CELL_N)
    return p.parse_args()


def _spearman(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) < 3 or len(xs) != len(ys):
        return None

    def ranks(vals: list[float]) -> list[float]:
        order = sorted(range(len(vals)), key=lambda i: vals[i])
        r = [0.0] * len(vals)
        i = 0
        while i < len(order):
            j = i
            while j + 1 < len(order) and vals[order[j + 1]] == vals[order[i]]:
                j += 1
            avg = (i + j) / 2.0 + 1.0
            for k in range(i, j + 1):
                r[order[k]] = avg
            i = j + 1
        return r

    rx, ry = ranks(xs), ranks(ys)
    mx, my = statistics.mean(rx), statistics.mean(ry)
    num = sum((a - mx) * (b - my) for a, b in zip(rx, ry))
    denx = math.sqrt(sum((a - mx) ** 2 for a in rx))
    deny = math.sqrt(sum((b - my) ** 2 for b in ry))
    if denx == 0 or deny == 0:
        return None
    return num / (denx * deny)


def _mae(xs: list[float], ys: list[float]) -> float | None:
    if not xs:
        return None
    return statistics.mean(abs(a - b) for a, b in zip(xs, ys))


def _composite(scores: dict[str, float]) -> float | None:
    try:
        q = sum(Q_WEIGHTS[d] * float(scores[d]) for d in Q_WEIGHTS)
        ef = 0.70 + 0.03 * float(scores["evidence_strength"])
        return min(10.0, q * ef)
    except (KeyError, TypeError, ValueError):
        return None


def _has_full_dims(scores: dict[str, Any]) -> bool:
    return all(scores.get(d) is not None for d in QUALITY_DIMS)


def _load_humans(labeller: str, label_round: str) -> list[dict[str, Any]]:
    with connect() as conn:
        rows = conn.execute(
            """
            SELECT *
            FROM paper_intelligence.golden_human_scores
            WHERE labeller = %s AND label_round = %s
            ORDER BY paper_id
            """,
            (labeller, label_round),
        ).fetchall()
    out = []
    for r in rows:
        d = dict(r)
        scores = {}
        for dim in QUALITY_DIMS:
            key = f"h_{dim}"
            if d.get(key) is not None:
                scores[dim] = float(d[key])
        # Optional seats kept for excel only
        for dim in ("tech_relevance", "product_relevance", "ai_relevance"):
            key = f"h_{dim}"
            if d.get(key) is not None:
                scores[dim] = float(d[key])
        final = float(d["h_final_score"]) if d.get("h_final_score") is not None else None
        if final is not None:
            composite = final
        else:
            composite = _composite({k: scores[k] for k in QUALITY_DIMS if k in scores})
        out.append(
            {
                "paper_id": int(d["paper_id"]),
                "arxiv_id": d.get("arxiv_id"),
                "sample_stratum": d.get("sample_stratum") or "unknown",
                "verdict": d.get("h_newsletter_verdict"),
                "scores": scores,
                "composite": composite,
                "h_final_score": final,
                "title": None,
            }
        )
    return out


def _attach_titles(humans: list[dict[str, Any]]) -> None:
    ids = [h["paper_id"] for h in humans]
    with connect() as conn:
        rows = conn.execute(
            "SELECT paper_id, title FROM paper_intelligence.papers WHERE paper_id = ANY(%s)",
            (ids,),
        ).fetchall()
    titles = {int(r["paper_id"]): r["title"] for r in rows}
    for h in humans:
        h["title"] = titles.get(h["paper_id"])


def _load_terra(ids: list[int]) -> dict[int, dict[str, Any]]:
    """Latest Terra quality scores, excluding G3b evaluation runs."""
    out: dict[int, dict[str, Any]] = {}
    excluded = _excluded_terra_run_ids()
    with connect() as conn:
        if excluded:
            rows = conn.execute(
                """
                SELECT DISTINCT ON (q.content_item_id)
                  q.content_item_id, q.result_json, q.run_id
                FROM paper_intelligence.paper_classification_results q
                WHERE q.task_type = 'quality'
                  AND q.model LIKE %s
                  AND q.content_item_id = ANY(%s)
                  AND (q.run_id IS NULL OR NOT (q.run_id = ANY(%s::uuid[])))
                ORDER BY q.content_item_id, q.created_at DESC
                """,
                ("%terra%", ids, excluded),
            ).fetchall()
        else:
            rows = conn.execute(
                """
                SELECT DISTINCT ON (q.content_item_id)
                  q.content_item_id, q.result_json, q.run_id
                FROM paper_intelligence.paper_classification_results q
                WHERE q.task_type = 'quality'
                  AND q.model LIKE %s
                  AND q.content_item_id = ANY(%s)
                ORDER BY q.content_item_id, q.created_at DESC
                """,
                ("%terra%", ids),
            ).fetchall()
    for r in rows:
        rj = r["result_json"] or {}
        if isinstance(rj, str):
            rj = json.loads(rj)
        dims = {}
        ok = True
        for d in QUALITY_DIMS:
            if rj.get(d) is None:
                ok = False
                break
            dims[d] = float(rj[d])
        if not ok:
            continue
        comp = rj.get("composite") or {}
        # Match G2: prefer stored quality (weighted sum); fall back to local composite.
        if isinstance(comp, dict) and comp.get("quality") is not None:
            quality = float(comp["quality"])
        else:
            quality = _composite(dims)
        out[int(r["content_item_id"])] = {"scores": dims, "composite": quality}
    return out


def _load_jev(path: Path) -> dict[int, dict[str, Any]]:
    if not path.exists():
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    out: dict[int, dict[str, Any]] = {}
    for k, v in (raw.get("papers") or {}).items():
        dims = dict(v.get("dims") or {})
        if not all(dims.get(d) is not None for d in QUALITY_DIMS):
            continue
        try:
            scores = {d: float(dims[d]) for d in QUALITY_DIMS}
        except (KeyError, TypeError, ValueError):
            continue
        comp = v.get("composite")
        if comp is None:
            comp = _composite(scores)
        out[int(k)] = {"scores": scores, "composite": float(comp) if comp is not None else None}
    return out


def _pair_stats(
    humans: list[dict[str, Any]],
    engine: dict[int, dict[str, Any]],
    dim: str,
) -> dict[str, Any]:
    xs, ys = [], []
    for h in humans:
        if dim not in h["scores"]:
            continue
        e = engine.get(h["paper_id"])
        if not e or dim not in e["scores"]:
            continue
        xs.append(float(h["scores"][dim]))
        ys.append(float(e["scores"][dim]))
    return {
        "n": len(xs),
        "spearman": _spearman(xs, ys) if len(xs) >= MIN_CELL_N else None,
        "mae": _mae(xs, ys) if len(xs) >= MIN_CELL_N else None,
        "suppressed": len(xs) < MIN_CELL_N,
    }


def _composite_stats(
    humans: list[dict[str, Any]], engine: dict[int, dict[str, Any]]
) -> dict[str, Any]:
    xs, ys = [], []
    for h in humans:
        if h.get("composite") is None:
            continue
        e = engine.get(h["paper_id"])
        if not e or e.get("composite") is None:
            continue
        xs.append(float(h["composite"]))
        ys.append(float(e["composite"]))
    return {
        "n": len(xs),
        "spearman": _spearman(xs, ys) if len(xs) >= MIN_CELL_N else None,
        "mae": _mae(xs, ys) if len(xs) >= MIN_CELL_N else None,
        "suppressed": len(xs) < MIN_CELL_N,
    }


def _verdict_recall(
    humans: list[dict[str, Any]],
    engine: dict[int, dict[str, Any]],
    *,
    top_n: int,
) -> dict[str, Any]:
    winners = [h for h in humans if h.get("verdict") == "winner_material"]
    scored = []
    for h in humans:
        e = engine.get(h["paper_id"])
        if e and e.get("composite") is not None:
            scored.append((float(e["composite"]), h["paper_id"]))
    scored.sort(key=lambda t: (-t[0], t[1]))
    top_ids = {cid for _, cid in scored[:top_n]}
    hit = sum(1 for w in winners if w["paper_id"] in top_ids)
    return {
        "winner_material_n": len(winners),
        "top_n": top_n,
        "hits": hit,
        "recall": (hit / len(winners)) if winners else None,
        "engine_scored_n": len(scored),
    }


def _closer_engine(stats_by_engine: dict[str, dict[str, Any]]) -> str | None:
    best = None
    best_key = None
    for eng, st in stats_by_engine.items():
        if st.get("suppressed"):
            continue
        sp = st.get("spearman")
        if sp is None or st.get("n", 0) < MIN_CELL_N:
            continue
        mae = st.get("mae") if st.get("mae") is not None else 99.0
        key = (sp, -mae)
        if best_key is None or key > best_key:
            best_key = key
            best = eng
    return best


def _fmt(x: float | None, nd: int = 3) -> str:
    return "n/a" if x is None else f"{x:.{nd}f}"


def _cell(st: dict[str, Any]) -> str:
    if st.get("suppressed") or st.get("n", 0) < MIN_CELL_N:
        return f"suppressed (n={st['n']})"
    return f"{_fmt(st['spearman'])} / {_fmt(st['mae'], 2)} / {st['n']}"


def _build_block(
    humans: list[dict[str, Any]], engines: dict[str, dict[int, dict[str, Any]]]
) -> dict[str, Any]:
    block: dict[str, Any] = {
        "n": len(humans),
        "dims": {},
        "composite": {},
        "verdict_recall": {},
    }
    for dim in QUALITY_DIMS:
        by_eng = {e: _pair_stats(humans, engines[e], dim) for e in ENGINES}
        block["dims"][dim] = {**by_eng, "closer_to_human": _closer_engine(by_eng)}
    by_eng_c = {e: _composite_stats(humans, engines[e]) for e in ENGINES}
    block["composite"] = {**by_eng_c, "closer_to_human": _closer_engine(by_eng_c)}
    for e in ENGINES:
        block["verdict_recall"][e] = {
            "top_20": _verdict_recall(humans, engines[e], top_n=20),
            "top_50": _verdict_recall(humans, engines[e], top_n=50),
        }
    return block


def main() -> int:
    global MIN_CELL_N
    args = _args()
    MIN_CELL_N = args.min_cell_n

    humans = _load_humans(args.labeller, args.label_round)
    if len(humans) < 150:
        print(
            f"CANNOT BE DETERMINED: only {len(humans)} labelled rows (need ≥150)",
            file=sys.stderr,
        )
    _attach_titles(humans)
    ids = [h["paper_id"] for h in humans]
    id_set = set(ids)

    terra_all = _load_terra(ids)
    jev1_all_raw = _load_jev(JEV_V001)
    jev2_all_raw = _load_jev(JEV_V002)
    # Restrict Jev caches to golden set only (caches contain thousands of extras)
    jev1_all = {k: v for k, v in jev1_all_raw.items() if k in id_set}
    jev2_all = {k: v for k, v in jev2_all_raw.items() if k in id_set}

    engines_full = {"terra": terra_all, "jev_v001": jev1_all, "jev_v002": jev2_all}

    # Intersection: papers with full 6 dims + composite on all three engines
    inter_ids = set()
    for h in humans:
        cid = h["paper_id"]
        ok = True
        for eng in ENGINES:
            e = engines_full[eng].get(cid)
            if not e or not _has_full_dims(e["scores"]) or e.get("composite") is None:
                ok = False
                break
        if ok and _has_full_dims(h["scores"]) and h.get("composite") is not None:
            inter_ids.add(cid)

    humans_inter = [h for h in humans if h["paper_id"] in inter_ids]
    engines_inter = {
        e: {cid: engines_full[e][cid] for cid in inter_ids} for e in ENGINES
    }

    overall_inter = _build_block(humans_inter, engines_inter)
    overall_full = _build_block(humans, engines_full)

    # Per-stratum on intersection only; suppress strata with n < MIN_CELL_N
    by_stratum: dict[str, Any] = {}
    strata_suppressed: list[dict[str, Any]] = []
    strata = sorted({h["sample_stratum"] for h in humans_inter})
    for stratum in strata:
        subset = [h for h in humans_inter if h["sample_stratum"] == stratum]
        if len(subset) < MIN_CELL_N:
            strata_suppressed.append(
                {"stratum": stratum, "n": len(subset), "reason": f"n < {MIN_CELL_N}"}
            )
            continue
        by_stratum[stratum] = _build_block(subset, engines_inter)

    # Also note full-set strata that never enter intersection (flagged/screen_failed)
    for stratum in sorted({h["sample_stratum"] for h in humans}):
        n_full = sum(1 for h in humans if h["sample_stratum"] == stratum)
        n_inter = sum(1 for h in humans_inter if h["sample_stratum"] == stratum)
        if n_inter < MIN_CELL_N and stratum not in {s["stratum"] for s in strata_suppressed}:
            strata_suppressed.append(
                {
                    "stratum": stratum,
                    "n": n_inter,
                    "n_labelled": n_full,
                    "reason": f"intersection n < {MIN_CELL_N}",
                }
            )

    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "phase": "G3a",
        "labeller": args.labeller,
        "label_round": args.label_round,
        "n_labelled": len(humans),
        "n_intersection": len(humans_inter),
        "min_cell_n": MIN_CELL_N,
        "gate_min_rows": 150,
        "cannot_be_determined": len(humans_inter) < 150,
        "engine_coverage_full": {e: len(engines_full[e]) for e in ENGINES},
        "engine_coverage_intersection": {e: len(engines_inter[e]) for e in ENGINES},
        "notes": [
            "ai_relevance REMOVED from all comparison tables. It was backfilled from "
            "pipeline screen scores for this labelling round; a Spearman of 1.000 vs "
            "Terra was circular (not independent human agreement) and would be misread.",
            "Headline metrics use the INTERSECTION of papers scored by all three engines "
            f"(Terra + Jev v001 + Jev v002) with full six quality dimensions (n={len(humans_inter)}).",
            f"Any stratum or cell with n < {MIN_CELL_N} is suppressed (noise, not a finding).",
            "Secondary table reports per-engine full coverage on the labelled set, with n per cell.",
            "Seat scores (tech_relevance / product_relevance) remain unscored by engines — omitted.",
            "Composite comparison uses human h_final_score vs engine quality composite.",
            "G3b evaluation run_ids listed in reports/golden/g3b_run.json are excluded "
            "from Terra loads here (gate-drop eval must not pollute current quality).",
        ],
        "overall_intersection": overall_inter,
        "overall_full_coverage": overall_full,
        "by_stratum_intersection": by_stratum,
        "strata_suppressed": strata_suppressed,
        "verdict_counts_labelled": dict(Counter(h.get("verdict") for h in humans)),
        "verdict_counts_intersection": dict(
            Counter(h.get("verdict") for h in humans_inter)
        ),
    }

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    JSON_OUT.write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")

    # ---- Markdown ----
    lines = [
        "# Golden engine comparison (Phase G3a)",
        "",
        f"**Generated:** {summary['generated_at']}  ",
        f"**Labels:** labeller=`{args.labeller}` round=`{args.label_round}` "
        f"n_labelled=**{len(humans)}** · n_intersection=**{len(humans_inter)}**  ",
        f"**Verdict mix (intersection):** {summary['verdict_counts_intersection']}  ",
        f"**Engine coverage (full / intersection):** "
        f"{summary['engine_coverage_full']} / {summary['engine_coverage_intersection']}  ",
        f"**Gate (intersection ≥150):** "
        f"{'OK' if not summary['cannot_be_determined'] else 'CANNOT BE DETERMINED'}",
        "",
        "## Notes",
        "",
    ]
    for n in summary["notes"]:
        lines.append(f"- {n}")

    lines += [
        "",
        f"## Headline — intersection (n={len(humans_inter)}, all three engines)",
        "",
        "Per dimension: Spearman / MAE / n",
        "",
        "| Dimension | Terra | Jev v001 | Jev v002 | Closer to human |",
        "|---|---:|---:|---:|---|",
    ]
    for dim in QUALITY_DIMS:
        block = overall_inter["dims"][dim]
        cells = [_cell(block[e]) for e in ENGINES]
        closer = block["closer_to_human"] or "n/a"
        lines.append(f"| {dim} | {cells[0]} | {cells[1]} | {cells[2]} | **{closer}** |")
    cc = overall_inter["composite"]
    cells = [_cell(cc[e]) for e in ENGINES]
    lines.append(
        f"| **composite (h_final vs engine)** | {cells[0]} | {cells[1]} | {cells[2]} | "
        f"**{cc['closer_to_human'] or 'n/a'}** |"
    )

    lines += [
        "",
        "### Verdict recall on intersection (`winner_material` in engine top-K)",
        "",
        "| Engine | winners | top-20 hits | recall@20 | top-50 hits | recall@50 |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for e in ENGINES:
        t20 = overall_inter["verdict_recall"][e]["top_20"]
        t50 = overall_inter["verdict_recall"][e]["top_50"]
        lines.append(
            f"| {e} | {t50['winner_material_n']} | {t20['hits']} | {_fmt(t20['recall'], 3)} | "
            f"{t50['hits']} | {_fmt(t50['recall'], 3)} |"
        )

    lines += [
        "",
        "## Secondary — per-engine full coverage (n varies by cell)",
        "",
        "Same labelled set; each engine uses whatever papers it scored. "
        f"Cells with n < {MIN_CELL_N} suppressed.",
        "",
        "| Dimension | Terra | Jev v001 | Jev v002 |",
        "|---|---:|---:|---:|",
    ]
    for dim in QUALITY_DIMS:
        block = overall_full["dims"][dim]
        lines.append(
            f"| {dim} | {_cell(block['terra'])} | {_cell(block['jev_v001'])} | "
            f"{_cell(block['jev_v002'])} |"
        )
    cc_f = overall_full["composite"]
    lines.append(
        f"| **composite** | {_cell(cc_f['terra'])} | {_cell(cc_f['jev_v001'])} | "
        f"{_cell(cc_f['jev_v002'])} |"
    )
    lines += [
        "",
        "| Engine | winners (labelled) | recall@20 | recall@50 | engine_scored_n |",
        "|---|---:|---:|---:|---:|",
    ]
    for e in ENGINES:
        t20 = overall_full["verdict_recall"][e]["top_20"]
        t50 = overall_full["verdict_recall"][e]["top_50"]
        lines.append(
            f"| {e} | {t50['winner_material_n']} | {_fmt(t20['recall'], 3)} | "
            f"{_fmt(t50['recall'], 3)} | {t50['engine_scored_n']} |"
        )

    lines += ["", f"## Per stratum (intersection only; n ≥ {MIN_CELL_N})", ""]
    if not by_stratum:
        lines.append("_No strata meet the n ≥ 25 threshold._")
        lines.append("")
    for stratum, block in by_stratum.items():
        lines += [
            f"### `{stratum}` (n={block['n']})",
            "",
            "| Dimension | Terra sp | Jev v001 sp | Jev v002 sp | Closer |",
            "|---|---:|---:|---:|---|",
        ]
        for dim in QUALITY_DIMS:
            b = block["dims"][dim]
            lines.append(
                f"| {dim} | {_fmt(b['terra']['spearman'])} | "
                f"{_fmt(b['jev_v001']['spearman'])} | {_fmt(b['jev_v002']['spearman'])} | "
                f"**{b['closer_to_human'] or 'n/a'}** |"
            )
        b = block["composite"]
        lines.append(
            f"| composite | {_fmt(b['terra']['spearman'])} | "
            f"{_fmt(b['jev_v001']['spearman'])} | {_fmt(b['jev_v002']['spearman'])} | "
            f"**{b['closer_to_human'] or 'n/a'}** |"
        )
        lines.append("")

    if strata_suppressed:
        lines += [
            "### Suppressed strata",
            "",
            "| Stratum | n (intersection) | Reason |",
            "|---|---:|---|",
        ]
        for s in strata_suppressed:
            lines.append(
                f"| `{s['stratum']}` | {s['n']} | {s['reason']} |"
            )
        lines.append("")

    lines += [
        "## STOP",
        "",
        "- No model / prompt / policy / threshold / weight changes from this run.",
        f"- Baseline recorded at `{BASELINE_OUT.relative_to(ROOT)}` from **intersection** numbers.",
        f"- Artifacts: `{MD_OUT.relative_to(ROOT)}`, `{JSON_OUT.relative_to(ROOT)}`, "
        f"`{XLSX_OUT.relative_to(ROOT)}`",
        "",
    ]
    MD_OUT.write_text("\n".join(lines), encoding="utf-8")

    # ---- Baseline from intersection ----
    base_lines = [
        "# Golden baseline v1",
        "",
        f"**Recorded:** {summary['generated_at']}  ",
        f"**Labels:** `{args.labeller}` / `{args.label_round}` · "
        f"n_labelled={len(humans)} · **n_intersection={len(humans_inter)}**  ",
        "",
        "**ai_relevance removed:** backfilled from screen scores; Spearman 1.000 vs Terra "
        "was circular, not human agreement. Do not restore without independent human labels.",
        "",
        "Verifier gate (scoring changes) — measured on the **3-engine intersection**:",
        "- Per-dimension Spearman vs human must not fall more than **0.05** below these values.",
        "- Verdict recall@50 must not fall below the value below.",
        "- Intersection n < 150 → Cannot be determined.",
        f"- Cells / strata with n < {MIN_CELL_N} are not gate inputs.",
        "",
        f"## Per-dimension Spearman (intersection n={len(humans_inter)})",
        "",
        "| Dimension | Terra | Jev v001 | Jev v002 |",
        "|---|---:|---:|---:|",
    ]
    for dim in QUALITY_DIMS:
        b = overall_inter["dims"][dim]
        base_lines.append(
            f"| {dim} | {_fmt(b['terra']['spearman'])} | "
            f"{_fmt(b['jev_v001']['spearman'])} | {_fmt(b['jev_v002']['spearman'])} |"
        )
    base_lines += [
        "",
        f"Composite Spearman — Terra `{_fmt(overall_inter['composite']['terra']['spearman'])}` · "
        f"Jev v001 `{_fmt(overall_inter['composite']['jev_v001']['spearman'])}` · "
        f"Jev v002 `{_fmt(overall_inter['composite']['jev_v002']['spearman'])}`",
        "",
        "## Verdict recall@50 (intersection baseline)",
        "",
    ]
    for e in ENGINES:
        t50 = overall_inter["verdict_recall"][e]["top_50"]
        base_lines.append(
            f"- **{e}:** {_fmt(t50['recall'], 3)} "
            f"({t50['hits']}/{t50['winner_material_n']} winner_material in top 50)"
        )
    base_lines += [
        "",
        f"Source: `{MD_OUT.relative_to(ROOT)}`",
        "",
    ]
    BASELINE_OUT.write_text("\n".join(base_lines), encoding="utf-8")

    # ---- Excel ----
    wb = Workbook()
    ws2 = wb.active
    ws2.title = "engine_summary"
    ws2["A1"] = "Scope"
    ws2["B1"] = "Dimension"
    ws2["C1"] = "Terra Spearman"
    ws2["D1"] = "Jev v001 Spearman"
    ws2["E1"] = "Jev v002 Spearman"
    ws2["F1"] = "Closer"
    ws2["G1"] = "n"
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    r = 2
    for scope, block in (
        ("intersection", overall_inter),
        ("full_coverage", overall_full),
    ):
        for dim in QUALITY_DIMS:
            b = block["dims"][dim]
            ws2.cell(r, 1, scope)
            ws2.cell(r, 2, dim)
            ws2.cell(r, 3, b["terra"]["spearman"])
            ws2.cell(r, 4, b["jev_v001"]["spearman"])
            ws2.cell(r, 5, b["jev_v002"]["spearman"])
            ws2.cell(r, 6, b["closer_to_human"])
            ws2.cell(r, 7, b["terra"]["n"])
            r += 1
        b = block["composite"]
        ws2.cell(r, 1, scope)
        ws2.cell(r, 2, "composite")
        ws2.cell(r, 3, b["terra"]["spearman"])
        ws2.cell(r, 4, b["jev_v001"]["spearman"])
        ws2.cell(r, 5, b["jev_v002"]["spearman"])
        ws2.cell(r, 6, b["closer_to_human"])
        ws2.cell(r, 7, b["terra"]["n"])
        r += 1

    ws = wb.create_sheet("comparison")
    headers = [
        "paper_id",
        "arxiv_id",
        "sample_stratum",
        "in_intersection",
        "title",
        "h_verdict",
        "h_final",
        *[f"h_{d}" for d in QUALITY_DIMS],
        "terra_composite",
        *[f"terra_{d}" for d in QUALITY_DIMS],
        "jev_v001_composite",
        *[f"jev_v001_{d}" for d in QUALITY_DIMS],
        "jev_v002_composite",
        *[f"jev_v002_{d}" for d in QUALITY_DIMS],
        "delta_terra_final",
        "delta_jev_v001_final",
        "delta_jev_v002_final",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(1, col, h).font = Font(bold=True)
    for i, h in enumerate(humans, start=2):
        cid = h["paper_id"]
        t = terra_all.get(cid) or {"scores": {}, "composite": None}
        j1 = jev1_all.get(cid) or {"scores": {}, "composite": None}
        j2 = jev2_all.get(cid) or {"scores": {}, "composite": None}
        row = {
            "paper_id": cid,
            "arxiv_id": h.get("arxiv_id"),
            "sample_stratum": h["sample_stratum"],
            "in_intersection": cid in inter_ids,
            "title": h.get("title"),
            "h_verdict": h.get("verdict"),
            "h_final": h.get("h_final_score"),
            **{f"h_{d}": h["scores"].get(d) for d in QUALITY_DIMS},
            "terra_composite": t.get("composite"),
            **{f"terra_{d}": t["scores"].get(d) for d in QUALITY_DIMS},
            "jev_v001_composite": j1.get("composite"),
            **{f"jev_v001_{d}": j1["scores"].get(d) for d in QUALITY_DIMS},
            "jev_v002_composite": j2.get("composite"),
            **{f"jev_v002_{d}": j2["scores"].get(d) for d in QUALITY_DIMS},
        }
        if h.get("h_final_score") is not None:
            if t.get("composite") is not None:
                row["delta_terra_final"] = float(t["composite"]) - float(h["h_final_score"])
            if j1.get("composite") is not None:
                row["delta_jev_v001_final"] = float(j1["composite"]) - float(h["h_final_score"])
            if j2.get("composite") is not None:
                row["delta_jev_v002_final"] = float(j2["composite"]) - float(h["h_final_score"])
        for col, name in enumerate(headers, 1):
            ws.cell(i, col, row.get(name))

    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_OUT)

    print(f"wrote {MD_OUT}", flush=True)
    print(f"wrote {JSON_OUT}", flush=True)
    print(f"wrote {BASELINE_OUT}", flush=True)
    print(f"wrote {XLSX_OUT}", flush=True)
    print(
        json.dumps(
            {
                "n_labelled": len(humans),
                "n_intersection": len(humans_inter),
                "closer_composite_intersection": overall_inter["composite"]["closer_to_human"],
                "recall50_intersection": {
                    e: overall_inter["verdict_recall"][e]["top_50"]["recall"] for e in ENGINES
                },
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
